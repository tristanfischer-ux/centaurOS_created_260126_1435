#!/usr/bin/env python3
"""Build the marketplace-offtake screen in the Version 5 business-plan shape.

The shape is the 5 October 2026 Version 5 summary: Heading 1 in navy, and a
Side Heading frame in the left margin in navy (000080).
"""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

TEMPLATE = Path("/tmp/bp/plan.docx")
OUT = Path("/workspace/Business/opportunity-screens/2026-10-06-selling-into-a-marketplace.docx")


def blocks() -> list[tuple]:
    return [
        ("T", "Fractional Forge Limited"),
        ("T", "Selling into a marketplace"),
        ("T", "Five ways to be paid without a sales meeting"),
        ("T", "The opportunity, the product, the first customers, the costs and financial results, the risks, the team and funding needed, and an honest verdict"),
        ("T", "Draft of 6 October 2026"),
        ("T", "Draft. Confidential."),
        ("T", "Tristan Fischer, Fractional Forge Limited"),
        ("C", "Contents"),
        ("C", "1   The opportunity and the customers"),
        ("C", "2   The product"),
        ("C", "3   The first customers"),
        ("C", "4   Costs and the financial results"),
        ("C", "5   What one paid job costs, and which part is ours"),
        ("C", "6   Why the price compares as it does"),
        ("C", "7   Risks, permits and compliance"),
        ("C", "8   Team and funding needed"),
        ("C", "9   Viability verdict"),
        ("C", "10  What is missing and needs more work"),
        ("H", "The opportunity and the customers"),
        ("S", "The buyer is a market, not a meeting.",
         "The search was for a place that already pays for a standard unit, the way a generator is paid for electricity it delivers to the grid. A pitch to one named company was not enough. A spacecraft, a field crew, and a hire before the first unit were ruled out. Where the unit still has to be chosen by a person, the plan says so."),
        ("S", "Scores answer whether one person can start.",
         "Each idea was scored from one to five on whether the market clears without a pitch, whether one person can make the unit, how little capital the first sale needs, how fast cash can arrive, how much current models help, whether the market is shown to pay, and whether a United Kingdom person can register. Weights are 1.5, 1.3, 1.3, 1.2, 1.2, 1.0 and 1.0. A high score is not the same thing as the purest auction. The pure auctions have a megawatt floor."),
        ("S", "Twenty mechanisms are real. Ten are written longer. Five are the plan.",
         "Four were looked at and left out before the twenty. A residential-bandwidth proxy was left out because the buyer is often hiding the traffic. Puro.earth's own order book was left out of the five because the company says it discontinued that marketplace in 2022. A capacity-market seat and a carbon exchange are in the twenty, then cut, because you cannot post a unit you do not already hold."),
        ("S", "1. Piclo Flex. Weighted 3.38.",
         "Distribution operators buy flexibility on Piclo. National Grid's distribution arm runs week-ahead rounds, pay as bid, and from 1 April 2026 a day-ahead path on Piclo with gate closure at 11.30 the day before. Northern Powergrid's November 2026 specification asks for at least 10 kilowatts, half-hourly metering, and expects to procure up to 1 megawatt in that round. You upload an asset. You do not book a meeting. You do need the asset, inside a zone that is actually buying."),
        ("S", "2. A standing export tariff. Weighted 3.24.",
         "Octopus's export-only Smart Export Guarantee pays 4.1 pence per kilowatt-hour, and you may keep another supplier for the power you buy. Its export page states 12 pence per kilowatt-hour for a solar export tariff aimed at businesses and charities. A secondary page dated 25 September 2026 also states 12 pence for Outgoing Octopus where you buy your import from Octopus. Agile export follows half-hourly wholesale prices. Once the export meter is registered, the cheque does not need a customer. It needs panels, a smart meter, and an MCS or Flexi-Orb certificate, or Octopus's paid certification route."),
        ("S", "3. NESO dynamic response. Weighted 2.69.",
         "Dynamic Containment, Dynamic Moderation and Dynamic Regulation are day-ahead auctions. NESO's own note says a minimum response capacity of 1 megawatt per unit, sell orders in whole megawatts, and payment for availability. Aggregation is allowed inside a grid supply point group. A proposal to cut a different service, static response, to 0.1 megawatt is a 2026 reform, not the rule for these three. This is the grid. It is not a solo start."),
        ("S", "4. The Capacity Market. Weighted 2.55.",
         "A demand-side unit must show a proven capacity of at least 1 megawatt. Agreements awarded after the Capacity Market (Amendment) (No. 4) Rules 2026 must also show at least half of the obligation. This is an annual auction with credit cover, not a listing you open on a Tuesday."),
        ("S", "5. Vast.ai GPU hours. Weighted 3.59.",
         "You list a machine. Renters take hours at the price you set. Vast's own article illustrates 0.30 to 0.60 dollars per GPU-hour for an RTX 5090, and says earnings cannot be guaranteed. Verification wants on the order of 500 megabits per second and a public IPv4 address. Payout starts at 20 dollars, and the host must be able to receive a business payment. This is the closest digital version of a spot market."),
        ("S", "6. Akash. Weighted 3.09.",
         "Akash is a compute marketplace where the provider sets the price and is paid in ACT for leases. The project's own page puts hardware from about 500 dollars if you buy it, and electricity on top. Settlement is not pounds in a bank. It stays in the twenty and not the ten."),
        ("S", "7. AWS Data Exchange. Weighted 4.31.",
         "You register as an AWS Marketplace seller, ask to be onboarded for data, and subscribers are billed by AWS. AWS says it pays you monthly, after it has collected, net of its fees. A 2023 secondary page states a 3 percent fee on public data listings. That percentage was not re-read on the current fee schedule, so it is not used as a price. Selling into Europe adds a know-your-customer step. The buyer still chooses the dataset. The invoice does not."),
        ("S", "8. Snowflake Marketplace. Weighted 3.45.",
         "Paid listings need a full Snowflake account, a provider profile, Stripe, and an approach to Marketplace Operations or a business-development partner before the listing is approved. The provider remains the seller of record. It is a real rail and a slow gate."),
        ("S", "9. RapidAPI. Weighted 4.56.",
         "You set the price of an API. From 15 November 2025 the marketplace fee is 25 percent, so a 100 dollar plan pays you 75 dollars before PayPal's own fee. Payout is by PayPal, at the end of the month after the charge. Developers still pick the API. Nobody has to take a call."),
        ("S", "10. Shopify App Store. Weighted 4.36.",
         "From 1 January 2025 a developer keeps 100 percent of the first 1 million dollars of lifetime gross app revenue, then 85 percent. There is a 2.9 percent processing fee and a 19 dollar registration. A merchant still chooses to install. The share is the most generous of the software rails in this list."),
        ("S", "11. Apple App Store. Weighted 4.22.",
         "The Small Business Program is a 15 percent commission when proceeds are at or under 1 million dollars in the prior year, and for a new developer. Above that, the standard commission returns. Review is a gate. The customer is still a person who taps Install."),
        ("S", "12. Google Play. Weighted 4.22.",
         "From 30 June 2026, in the United Kingdom, the United States and the European Economic Area, the service fee starts at 10 percent on the first 1 million dollars of annual earnings. If the charge uses Play billing, a further 5 percent billing fee applies in those places. Same shape as Apple, a different rate card."),
        ("S", "13. Adobe Stock. Weighted 4.33.",
         "A buyer licenses a file without meeting you. Generative images, vectors and video are accepted if you have the rights, you tick the generative box, and you do not put real people, brands or other artists into the prompt. The first payout needs 25 dollars and a wait of 45 days after the first sale, via PayPal, Payoneer or Skrill. The per-download royalty was not copied off the rate card in this pass, so no earnings figure is stated."),
        ("S", "14. Unity Asset Store. Weighted 3.99.",
         "The provider agreement says Unity pays 70 percent of the sales price, less refunds, transfer fees and tax. A game developer still chooses the asset. It is a shop, not a clearing price."),
        ("S", "15. Envato Elements. Weighted 3.99.",
         "Authors are paid from a share of subscriber revenue, with a stated allocation of half of net revenue from the base subscription price, and a quarter on enterprise subscriptions. Payout starts at 50 dollars. The buyer is a subscriber's download, not a purchase order. The catalogue is already full."),
        ("S", "16. Off-site biodiversity units. Weighted 2.29.",
         "England requires a 10 percent biodiversity net gain. A land manager can register a gain site. GOV.UK says the national register is not a matchmaking service and does not carry contact details. Buyers and sellers find each other privately. Statutory credits are a last resort sold by Natural England, not a bid you post. A 30-year legal agreement sits under every unit. This is a compliance duty. It is not a grid."),
        ("S", "17. Woodland Carbon Code. Weighted 2.18.",
         "Pending Issuance Units can be offered on an S&P Global request-for-information screen inside the UK Land Carbon Registry. There is no fee to look. The seller sets the price. The Code says these units cannot be listed on an exchange, and they are not guaranteed. You need a validated woodland."),
        ("S", "18. Nutrient credits. Weighted 2.14.",
         "Natural England sells credits to developers, continuously, in the Tees catchment at 2,700 pounds a credit and in Poole Harbour at 3,250 pounds, plus a non-refundable admin fee of 182.50 pounds plus VAT. One credit is one kilogram of nitrogen. Those are prices a developer pays the public scheme. They are not a price a landowner receives for posting an offer. The private supply side is still a contract."),
        ("S", "19. Xpansiv CBL. Weighted 2.71.",
         "CBL matches environmental commodities and settles product and cash, with spot settlement described as same day. You deposit units you already hold in a connected registry, then you offer them. It is a real exchange. It does not create the credit."),
        ("S", "20. Puro removal certificates. Weighted 1.71.",
         "A CORC is one tonne removed and stored. Puro says buyers deal with a supplier or a third-party channel, and that its own online marketplace stopped in 2022. An API can show unsold certificates to several sales channels at once. The unit is fungible. The factory is not a desk."),
        ("S", "What was cut, and why.",
         "The ten that stay for a longer look are Piclo, the export tariff, Vast, AWS Data Exchange, RapidAPI, NESO dynamic response, Shopify, Adobe Stock, the two phone stores taken together, and CBL. Akash goes because the payout is a token. Snowflake goes because of the approval gate. Unity and Envato go because they are shops and the founder has no edge in assets. Biodiversity, woodland and nutrients go because the public pages themselves say there is no order book, or because the published price is the price the state charges, not the price it pays. The Capacity Market goes because it is the same megawatt wall as NESO, with a slower auction."),
        ("H", "The product"),
        ("S", "Call Meter is an API with a published price.",
         "One narrow job, metered. A developer subscribes on Rapid. The first product should be something the founder can run on models he already uses: reading a planning or connection document and returning a structured answer, or turning a public earth-observation scene into a small JSON fact. The marketplace is the till. The product is the answer. A free tier exists so a developer can try a few calls, which is what Rapid itself recommends. The company does not take a briefing call."),
        ("S", "Scene List is a dataset with a subscription.",
         "One repeating file, on AWS Data Exchange. Not a platform. A monthly extract a modeller would otherwise build by hand: for example a change layer, or a table of sites, with the source scene named. AWS bills the subscriber. The founder does not. Private offers exist for a buyer who is already known. The public listing is the point of this plan. The 3 percent figure from a 2023 article is not used."),
        ("S", "Hour Desk is a machine on a spot board.",
         "One GPU host on Vast, priced by the founder, taken by whoever wants the hour. There is no customer success. There is uptime. Vast's illustration of 5090 rates is the company's own example. It is not a forecast. A machine that fails verification, or that sits on a slow line, does not get the rental. Do not buy a rack to test the sentence."),
        ("S", "Zone Bid is a kilowatt inside a zone that is buying.",
         "Piclo is the product only if a flexible asset sits in a published constraint zone. The asset can be a battery, a curtailable load, or generation that can turn up. Northern Powergrid's bar in the November 2026 paper is 10 kilowatts and half-hourly data. National Grid's day-ahead path pays as bid, against a ceiling the operator displays. If the map shows no competition on any site the founder can switch, Zone Bid does not exist yet. Aggregating other people's assets is a different business, and those people are customers."),
        ("S", "Export Cheque is a kilowatt-hour with a tariff, not a bid.",
         "The product is exported electricity under a published tariff. Octopus SEG at 4.1 pence is the clean export-only price. The 12 pence rates are different products and are not assumed to apply to a new roof. Agile is the one that behaves like a wholesale pass-through. Octopus's own page says the rate follows wholesale. A secondary page adds a floor of zero. That floor is not relied on. No annual yield is stated, because the roof does not exist in this plan."),
        ("S", "The other five in the ten are explanations, not products to build.",
         "NESO dynamic response is the auction the question was pointing at, and the 1 megawatt rule is why it is not started. Shopify, Apple and Google Play are real tills, and the user still chooses the app, so they fail the stricter test. Adobe Stock passes the stricter test on the buyer side and fails it on price, because the royalty per file was not retrieved. CBL is the till for a credit the founder does not have."),
        ("H", "The first customers"),
        ("S", "There is no first customer. There is a first listing.",
         "Call Meter's first counterparty is Rapid, then whichever developer turns on a paid plan. Scene List's first counterparty is AWS, then a subscriber. Hour Desk's first counterparty is a renter who never learns the founder's name. Zone Bid's counterparty is the distribution operator, if and only if the asset qualifies and the bid clears. Export Cheque's counterparty is the supplier that holds the export tariff, after the meter is registered."),
        ("S", "Do not start with the megawatt markets.",
         "NESO and the Capacity Market pay a clearing price, which is the right shape, and they require a megawatt and a registration stack. That is a project, not a first listing. CBL pays a matched price for a credit already in a registry. Issuing the credit is the whole business, and it is not this one."),
        ("S", "One channel each, and a public listing first.",
         "Call Meter ships one paid plan and one small free tier, then stops adding endpoints until someone has paid. Scene List ships one dataset and one price, then waits for a subscription before a second file. Hour Desk lists a machine the founder already has, or does not list. Zone Bid is a bid only after the Piclo map shows a live competition on a site the founder can control. Export Cheque is an application only for a system that already has a certificate."),
        ("H", "Costs and the financial results"),
        ("S", "No year-three profit is stated, because none has been observed.",
         "The figures below that are not on a public page are assumptions, labelled as such. They are not a forecast. The sourced facts are the fees and the floors. Where a United Kingdom price for our unit was not on a page, it is not invented."),
        ("S", "Call Meter: you keep 75 pounds of a 100 pound plan.",
         "That ratio is Rapid's published 25 percent fee, in force from 15 November 2025. PayPal's fee comes off after that. Payout is the month after the charge. Assumption, not a forecast: twenty subscribers at 20 pounds a month would be 300 pounds to the founder after the 25 percent, and before PayPal. Every term in that sum except the fee is an assumption. The cost to serve is model time. If the call needs a commercial satellite scene, the price has to carry the scene or the call is refused."),
        ("S", "Scene List: AWS pays after it collects.",
         "There is no sourced unit price for a dataset of this kind, so none is stated. Assumption: a subscription at 200 dollars a month. Ten subscribers would be 2,000 dollars a month before AWS's fee. The fee percentage is not taken from the 2023 secondary page. The know-your-customer step for European buyers is a delay, not a price. Gross margin stays high only while the file is computed from public scenes and public documents."),
        ("S", "Hour Desk: the company's example is not a budget.",
         "Vast's article says an RTX 5090 might earn 0.30 to 0.60 dollars per GPU-hour, and that a four-GPU rig at 80 percent utilisation might generate 700 to 1,400 dollars a month. Those sentences are the vendor's illustration. The same article says rates depend on hardware, uptime and demand. The docs say earnings cannot be guaranteed. A machine the founder does not own is not in the sum. Electricity is a real cost and was not priced here."),
        ("S", "Zone Bid and Export Cheque are site maths, not software maths.",
         "No Piclo clearing price is copied, because ceilings are per zone and were not retrieved for a named zone. A 10 kilowatt asset that never wins a bid earns nothing. Export at 4.1 pence per kilowatt-hour on 1,000 kilowatt-hours is 41 pounds. That arithmetic is the tariff, not a yield. Whether a roof exports 1,000 kilowatt-hours is not stated. At 4.1 pence, a new installation does not pay for itself inside this plan, and the plan does not pretend otherwise."),
        ("H", "What one paid job costs, and which part is ours"),
        ("S", "The founder's time is the cost, until a machine or a site enters.",
         "Call Meter and Scene List are time and model use. A commercial image is bought only if a named subscriber is paying more than the image. Hour Desk's cost is the machine, the power, and the line. Zone Bid's cost is the flexible asset and the metering. Export Cheque's cost is the generator and the certification. None of those capital items is spent in the first ninety days unless it is already owned."),
        ("S", "The platform's fee is the only third-party take on the desk products.",
         "Rapid takes 25 percent. AWS takes a marketplace fee that this document does not number. Apple's 15 percent and Google's 10 percent plus 5 percent billing fee, and Shopify's zero then 15 percent, are the comparisons. They are not the products being started. Vast does not publish a simple host commission in the pages read for this plan, so none is stated."),
        ("S", "Nothing here is a power station, and nothing is costed as one.",
         "There is no turbine, no habitat bank, and no credit inventory. The split is simple. The marketplace owns the till and the buyer. We own the answer, the file, the hour, or the kilowatt, and only the last two if the hardware is already there."),
        ("H", "Why the price compares as it does"),
        ("S", "A clearing price is honest. A shop price is a guess.",
         "Piclo and NESO publish a competition or an auction. The export tariff publishes a pence figure. Vast publishes a board of other people's asks. Rapid, AWS, Apple, Google and Shopify publish a fee, and you publish the price. Adobe publishes a royalty card that was not copied here. CBL publishes a market in credits we do not hold. Comparing a 20 pound API plan with 4.1 pence per kilowatt-hour is not a comparison of like for like. It is a comparison of what can be started."),
        ("S", "The phone stores and Shopify are the wrong shape, at a better fee.",
         "Shopify's zero fee on the first million dollars is a better commercial term than Rapid's 25 percent. It still needs a merchant to install an app. That is a customer, even when the shop window is the marketplace. The same is true of Apple at 15 percent and of Google Play at 10 percent plus a 5 percent billing fee in the United Kingdom when Play billing is used. They are in the ten so the fee is on the page. They are not in the five."),
        ("S", "AiDASH and the last plan are the warning, not the model.",
         "The previous screen killed a vegetation product because the buyer was an enterprise and the United States company was already in the market. This screen keeps only the rails where the till is the platform. A dataset that still needs a utility to sign a master service agreement is the old business, wearing a marketplace label. If Scene List only sells by private offer, it has failed the test and should be shut."),
        ("H", "Risks, permits and compliance"),
        ("S", "The listing can be live and still earn nothing.",
         "A Rapid plan with no subscriber is a zero. A dataset with no subscription is a zero. A GPU with no rental is a power bill. A Piclo asset outside a buying zone cannot bid. An export meter on a dark roof earns the tariff on nothing. None of these is a permit problem. All of them are the actual risk."),
        ("S", "The megawatt markets are a wall, not a ambition.",
         "Dynamic response at 1 megawatt, and a capacity-market test at 1 megawatt, are not a later phase of Call Meter. They are a different balance sheet. The static-response proposal at 0.1 megawatt is not in force for the dynamic products. Treating a proposal as a rule would be a false start."),
        ("S", "Rights, tax and the till.",
         "A generative file on Adobe Stock needs the rights the contributor agreement asks for, and the generative tick. An API that returns a fact about a person is a data-protection question and should not be the first endpoint. AWS selling into Europe requires the know-your-customer process they describe. Rapid pays PayPal, which a United Kingdom company must actually be able to receive. Vast requires a payout method that can take a business payment. None of this is an environmental permit. Grade the first listing against the till's own rules before the first file goes up."),
        ("H", "Team and funding needed"),
        ("S", "Tristan lists. Nobody is hired before a paid unit.",
         "One founder. No marketplace manager, no energy trader, and no field technician. A grid engineer is paid for an hour only if a real Piclo competition is on a site already controlled, and only to say whether the asset can perform. That hour is not a hire."),
        ("S", "No raise. The till does not need one.",
         "A raise would be spent finding customers, which is the cost this screen is trying not to have. Hardware for Hour Desk, Zone Bid or Export Cheque is bought only with money already available, and only after the listing rules are checked against that exact machine or site. Do not raise to buy a battery for a zone that has not published a competition."),
        ("S", "Control stays with the founder because there is no cap table event.",
         "There is no investor in this plan. If a raise is discussed later, it is after a paid subscription or a cleared bid, and it is a separate decision. This document does not authorise one."),
        ("H", "Viability verdict"),
        ("S", "Viable as a listing. Unproven until a stranger pays.",
         "The tills are real. Rapid's fee, AWS's disbursement, Vast's board, Piclo's competitions and Octopus's pence are sourced. The number of strangers who will pay for this founder's API, file, hour or kilowatt is not. The verdict is that two listings are worth ninety days, and that hardware is not bought to qualify for the other three."),
        ("S", "Start Call Meter and Scene List. Park the hardware.",
         "Both can be built with models and public inputs. Neither needs a panel, a GPU purchase, or a zone. Hour Desk starts only if a suitable machine and a fast public line are already in hand. Zone Bid starts only if the Piclo map shows a competition on a controllable site. Export Cheque starts only if a certified system already exists. Do not open a sixth listing in the same quarter. Do not build the Shopify, Apple or Google app in that quarter either."),
        ("S", "Three facts would most change the verdict.",
         "A Rapid listing that is allowed and still has no paid plan after ninety days. An AWS onboarding that refuses a one-person seller, or a fee schedule that takes most of a small subscription. A Piclo map with no competition anywhere the founder can switch a load. Any one of those, observed, demotes that idea. A beautiful auction at 1 megawatt does not promote NESO into the five."),
        ("H", "What is missing and needs more work"),
        ("S", "Every gap below has an owner and a next step.",
         "The owner is Tristan. Nothing below is delegated to a hire that does not exist."),
        ("S", "Eleven open items.",
         "One: publish one Rapid API with a paid plan in three weeks. Kill it if nobody has paid by day 90. Two: open the AWS Marketplace seller registration and read the current data-listing fee schedule before any price is quoted. The 3 percent figure stays unverified. Three: do not buy a GPU. If a machine is already owned, run Vast's verification and stop if the line or the address fails. Four: open Piclo, upload nothing, and record whether any live competition covers a site that can be controlled. Five: do not apply for an export tariff without an MCS or equivalent certificate already in hand. Six: confirm a United Kingdom company can receive Rapid's PayPal payout. Seven: do not treat Sunsave's averages as Octopus's prices. Eight: the Adobe royalty per download is still unread. Nine: static response at 0.1 megawatt is a proposal on a NESO procurement note, not a rule for dynamic containment. Ten: this pass used public pages, not an OpenRouter loop. Eleven: this document is a screen, not a decision to incorporate, and not a decision to buy hardware."),
    ]


def clear_body(doc: Document) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def add_navy_line(doc: Document, text: str, *, bold: bool = False, size: int = 11) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x80)
    run.font.size = Pt(size)
    run.font.name = "Arial"


def replace_header_footer(doc: Document) -> None:
    replacements = {
        "Version 5 summary business plan": "Selling into a marketplace",
        "5 October  2026": "6 October 2026",
        "5 October 2026": "6 October 2026",
    }
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            if part is None:
                continue
            for paragraph in part.paragraphs:
                for key, value in replacements.items():
                    if key in paragraph.text:
                        for run in paragraph.runs:
                            if key in run.text:
                                run.text = run.text.replace(key, value)


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(TEMPLATE, OUT)
    doc = Document(str(OUT))
    clear_body(doc)
    for block in blocks():
        kind = block[0]
        if kind == "T":
            paragraph = doc.add_paragraph(block[1], style="Title")
            if block[1].startswith("Draft"):
                paragraph.runs[0].bold = False
                paragraph.runs[0].font.size = Pt(12)
        elif kind == "C":
            add_navy_line(doc, block[1], bold=block[1] == "Contents", size=14 if block[1] == "Contents" else 11)
        elif kind == "H":
            doc.add_paragraph(block[1], style="Heading 1")
        elif kind == "S":
            doc.add_paragraph(block[1], style="Side Heading")
            doc.add_paragraph(block[2])
        else:
            raise ValueError(kind)
    replace_header_footer(doc)
    doc.save(str(OUT))
    print(OUT, OUT.stat().st_size)


if __name__ == "__main__":
    main()
